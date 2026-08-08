import json
import openpyxl
import emailhelper

from pathlib import Path
from logger import configura_logger


BASE_DIR = Path(__file__).resolve().parent
CONFIG_FILE = BASE_DIR / "config.json"


def carica_config():
    with open(CONFIG_FILE, "r", encoding="utf-8") as file:
        return json.load(file)


def carica_soglie(sheet_soglie, config):

    soglie = {}
    indexes = config["indexes_soglie"]

    for row in sheet_soglie.iter_rows(
        min_row=indexes["first_row"],
        values_only=True
    ):
        articolo = row[indexes["index_articolo"]]
        soglia = row[indexes["index_soglia"]]

        if articolo is None:
            continue

        if soglia is None:
            continue

        soglie[articolo] = soglia

    return soglie


def controlla_giacenza(config, logger):

    logger.info("Lettura file Excel")
    excel_file = config["excel_file_path"]

    workbook = openpyxl.load_workbook(
        excel_file,
        data_only=True
    )

    logger.info(
        f"File Excel caricato: {excel_file}"
    )

    sheet_magazzino = workbook[
        config["sheet_name_magazzino"]
    ]

    sheet_soglie = workbook[
        config["sheet_name_soglie"]
    ]

    indexes = config["indexes_magazzino"]

    # Carico tutte le soglie
    soglie = carica_soglie(sheet_soglie, config)

    logger.info(
        f"Soglie caricate: {len(soglie)}"
    )

    articoli_da_riordinare = []

    # Controllo il magazzino
    for row in sheet_magazzino.iter_rows(
        min_row=indexes["first_row"],
        values_only=True
    ):

        articolo = row[indexes["index_articolo"]]
        giacenza = row[indexes["index_giacenza"]]

        if not articolo:
            continue

        if giacenza is None:
            continue

        # Cerco la soglia dell'articolo
        soglia = soglie.get(articolo)

        if soglia is None:
            logger.warning(
                f"Nessuna soglia trovata per "
                f"'{articolo}'"
            )
            continue

        # Articolo sotto soglia
        if giacenza <= soglia:

            logger.info(
                f"Rifornimento necessario: "
                f"{articolo} "
                f"(giacenza={giacenza}, "
                f"soglia={soglia})"
            )

            articoli_da_riordinare.append({
                "articolo": articolo,
                "giacenza": giacenza,
                "soglia": soglia
            })

    workbook.close()

    logger.info(
        f"Articoli da riordinare: "
        f"{len(articoli_da_riordinare)}"
    )

    # Se ci sono articoli da riordinare
    if articoli_da_riordinare:

        try:

            emailhelper.invia_email(
                config,
                articoli_da_riordinare
            )

            logger.info(
                "Email di rifornimento inviata"
            )

        except Exception as e:

            logger.exception(
                "Errore durante l'invio della email"
            )

    else:

        logger.info(
            "Nessun articolo necessita di rifornimento"
        )


def main():

    config = carica_config()

    logger = configura_logger(config)

    logger.info("================================")
    logger.info("Controllo Magazzino")
    logger.info("================================")

    try:

        logger.info("Avvio controllo magazzino")

        controlla_giacenza(config, logger)

        logger.info("Controllo completato")

    except Exception:

        logger.exception(
            "Errore durante il controllo"
        )


if __name__ == "__main__":
    main()
